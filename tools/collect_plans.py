#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import json, os, re, sys, time, hashlib, mimetypes, subprocess, urllib.parse
from pathlib import Path

import requests
from bs4 import BeautifulSoup
import fitz
from docx import Document

BASE = Path(__file__).resolve().parent.parent
TARGETS = json.loads((BASE / 'tools' / 'plan_targets.json').read_text(encoding='utf-8'))
OUT = BASE / 'recovered_plans'
OUT.mkdir(exist_ok=True)
REPORT = []
UA = 'Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/131 Safari/537.36'
HEADERS = {'User-Agent': UA, 'Accept-Language': 'zh-CN,zh;q=0.9,en;q=0.5'}
SESSION = requests.Session()
SESSION.headers.update(HEADERS)
DOC_EXTS = ('.pdf', '.doc', '.docx', '.ofd')
BAD_TITLE = ('建议', '草案', '解读', '决议', '工作报告', '会议', '采购', '征求意见', '说明')


def log(msg):
    print(msg, flush=True)


def clean(s):
    return re.sub(r'\s+', '', s or '').replace('（', '(').replace('）', ')')


def safe(s):
    return re.sub(r'[\\/:*?"<>|]+', '_', s).strip()


def is_official(url):
    host = urllib.parse.urlparse(url).netloc.lower().split(':')[0]
    return host.endswith('.gov.cn') or host.endswith('gov.cn') or host in {
        'www.spb.gov.cn', 'ah.spb.gov.cn', 'dgrd.dg.gov.cn', 'www.mzrd.gov.cn',
        'www.qyrd.gov.cn', 'www.xmrd.gov.cn', 'hnrd.huainan.gov.cn',
        'ahczrd.gov.cn', 'www.ndwww.cn', 'www.cnbayarea.org.cn'
    }


def request(url, timeout=45):
    last = None
    for n in range(3):
        try:
            r = SESSION.get(url, timeout=timeout, allow_redirects=True, verify=True)
            if r.status_code < 500:
                return r
            last = RuntimeError(f'HTTP {r.status_code}')
        except Exception as e:
            last = e
        time.sleep(1.5 * (n + 1))
    raise last or RuntimeError('request failed')


def infer_ext(url, r, data):
    path = urllib.parse.urlparse(r.url or url).path.lower()
    for ext in DOC_EXTS:
        if path.endswith(ext):
            return ext
    cd = r.headers.get('content-disposition', '')
    m = re.search(r'filename\*?=(?:UTF-8\'\')?["\']?([^"\';]+)', cd, re.I)
    if m:
        name = urllib.parse.unquote(m.group(1))
        ext = Path(name).suffix.lower()
        if ext in DOC_EXTS:
            return ext
    ct = r.headers.get('content-type', '').lower()
    if 'pdf' in ct or data.startswith(b'%PDF'):
        return '.pdf'
    if 'wordprocessingml' in ct or data.startswith(b'PK\x03\x04'):
        return '.docx'
    if 'msword' in ct or data.startswith(bytes.fromhex('D0CF11E0')):
        return '.doc'
    if 'ofd' in ct:
        return '.ofd'
    return ''


def docx_text(path):
    try:
        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(' | '.join(c.text for c in row.cells))
        return '\n'.join(parts)
    except Exception:
        return ''


def pdf_stats(path):
    try:
        doc = fitz.open(path)
        pages = doc.page_count
        chunks = []
        for i in range(min(pages, 30)):
            try:
                chunks.append(doc[i].get_text('text'))
            except Exception:
                pass
        return pages, '\n'.join(chunks)
    except Exception:
        return 0, ''


def validate_document(path, area, direct_official=False):
    ext = path.suffix.lower()
    size = path.stat().st_size
    title_key = clean(area + '国民经济和社会发展第十五个五年规划纲要')
    generic_key = clean('国民经济和社会发展第十五个五年规划纲要')
    if size < 20000:
        return False, {'reason': 'file_too_small', 'size': size}
    if ext == '.pdf':
        pages, txt = pdf_stats(path)
        nt = clean(txt)
        title_hit = title_key in nt or (clean(area) in nt and generic_key in nt)
        if pages <= 0:
            return False, {'reason': 'invalid_pdf', 'size': size, 'pages': pages}
        if title_hit and (pages >= 5 or len(nt) >= 8000):
            return True, {'size': size, 'pages': pages, 'text_chars': len(nt), 'title_hit': True}
        # Scanned official attachments may have no text layer.
        if direct_official and pages >= 20 and size >= 150000:
            return True, {'size': size, 'pages': pages, 'text_chars': len(nt), 'title_hit': False, 'scan_or_image_pdf': True}
        return False, {'reason': 'title_or_length_check_failed', 'size': size, 'pages': pages, 'text_chars': len(nt), 'title_hit': title_hit}
    if ext == '.docx':
        txt = docx_text(path)
        nt = clean(txt)
        title_hit = title_key in nt or (clean(area) in nt and generic_key in nt)
        if title_hit and len(nt) >= 8000:
            return True, {'size': size, 'text_chars': len(nt), 'title_hit': True}
        return False, {'reason': 'docx_title_or_length_check_failed', 'size': size, 'text_chars': len(nt), 'title_hit': title_hit}
    if ext == '.doc':
        # Legacy .doc parsing is unreliable on runners; require a sizeable file from an official direct link.
        if direct_official and size >= 80000:
            return True, {'size': size, 'legacy_doc': True}
        return False, {'reason': 'legacy_doc_not_verified', 'size': size}
    if ext == '.ofd':
        if direct_official and size >= 80000:
            return True, {'size': size, 'ofd': True}
        return False, {'reason': 'ofd_not_verified', 'size': size}
    return False, {'reason': 'unknown_extension', 'size': size}


def try_download_document(url, outbase, area, direct_official=False):
    try:
        r = request(url, timeout=90)
    except Exception as e:
        return None, {'url': url, 'error': repr(e)}
    data = r.content
    ext = infer_ext(url, r, data)
    if not ext:
        return None, {'url': url, 'final_url': r.url, 'status': r.status_code, 'content_type': r.headers.get('content-type'), 'html': True}
    path = outbase.with_suffix(ext)
    path.write_bytes(data)
    ok, meta = validate_document(path, area, direct_official=direct_official and is_official(r.url or url))
    meta.update({'url': url, 'final_url': r.url, 'status': r.status_code, 'extension': ext})
    if ok:
        return path, meta
    try:
        path.unlink()
    except Exception:
        pass
    return None, meta


def attachment_candidates(html, base_url, area):
    soup = BeautifulSoup(html, 'html.parser')
    out = []
    seen = set()
    for a in soup.find_all('a'):
        href = a.get('href') or ''
        if not href:
            continue
        u = urllib.parse.urljoin(base_url, href)
        txt = a.get_text(' ', strip=True)
        combined = txt + ' ' + u
        lu = urllib.parse.urlparse(u).path.lower()
        if not (lu.endswith(DOC_EXTS) or 'download' in u.lower() or 'attachment' in u.lower() or 'fileurl=' in u.lower()):
            continue
        score = 0
        for kw, val in [('十五', 8), ('规划', 5), ('纲要', 7), (area, 4), ('文件下载', 3)]:
            if kw in combined:
                score += val
        if lu.endswith('.docx'):
            score += 4
        elif lu.endswith('.pdf'):
            score += 3
        elif lu.endswith('.doc'):
            score += 2
        if u not in seen:
            seen.add(u)
            out.append((score, u, txt))
    out.sort(reverse=True)
    return out[:15]


def page_text(html):
    soup = BeautifulSoup(html, 'html.parser')
    for t in soup(['script', 'style', 'noscript', 'svg']):
        t.decompose()
    return soup.get_text('\n', strip=True), (soup.title.get_text(' ', strip=True) if soup.title else '')


def looks_like_full_page(text, title, area):
    nt = clean(text)
    ntitle = clean(title)
    title_key = clean(area + '国民经济和社会发展第十五个五年规划纲要')
    generic_key = clean('国民经济和社会发展第十五个五年规划纲要')
    hit = title_key in nt or (clean(area) in nt and generic_key in nt)
    if not hit:
        return False
    if any(k in ntitle for k in BAD_TITLE) and len(nt) < 30000:
        return False
    chapters = len(re.findall(r'第[一二三四五六七八九十百]+章', text))
    return len(nt) >= 18000 or chapters >= 8


def chrome_binary():
    for x in ('google-chrome', 'google-chrome-stable', 'chromium', 'chromium-browser'):
        p = subprocess.run(['bash', '-lc', f'command -v {x}'], capture_output=True, text=True).stdout.strip()
        if p:
            return p
    return ''


CHROME = chrome_binary()


def print_page_pdf(url, dest):
    if not CHROME:
        return False, 'chrome_not_found'
    cmd = [CHROME, '--headless', '--no-sandbox', '--disable-gpu', '--disable-dev-shm-usage', '--hide-scrollbars',
           '--virtual-time-budget=12000', f'--print-to-pdf={dest}', url]
    try:
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=90)
    except Exception as e:
        return False, repr(e)
    return dest.exists() and dest.stat().st_size > 30000, (p.stderr or '')[-500:]


def markdown_to_docx(md, path, title, source_url):
    d = Document()
    d.add_heading(title, level=0)
    d.add_paragraph('来源：' + source_url)
    d.add_paragraph('说明：本文件根据上述政府官方网站公开全文整理，未改变正文内容。')
    for raw in md.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith('### '):
            d.add_heading(line[4:].strip(), level=3)
        elif line.startswith('## '):
            d.add_heading(line[3:].strip(), level=2)
        elif line.startswith('# '):
            d.add_heading(line[2:].strip(), level=1)
        else:
            d.add_paragraph(re.sub(r'\[(.*?)\]\([^)]*\)', r'\1', line))
    d.save(path)


def jina_fulltext(url, area, outbase):
    candidates = [
        'https://r.jina.ai/http://' + re.sub(r'^https?://', '', url),
        'https://r.jina.ai/https://' + re.sub(r'^https?://', '', url),
    ]
    for ju in candidates:
        try:
            r = request(ju, timeout=90)
        except Exception:
            continue
        md = r.text
        if looks_like_full_page(md, md[:500], area):
            path = outbase.with_suffix('.docx')
            markdown_to_docx(md, path, area + '国民经济和社会发展第十五个五年规划纲要（官方网页全文整理版）', url)
            ok, meta = validate_document(path, area, direct_official=False)
            if ok:
                meta.update({'url': url, 'jina_url': ju, 'method': 'jina_to_docx'})
                return path, meta
            try:
                path.unlink()
            except Exception:
                pass
    return None, {'url': url, 'method': 'jina_failed'}


def process_page(url, outbase, area, allow_full=True):
    notes = []
    try:
        r = request(url, timeout=60)
    except Exception as e:
        if allow_full and is_official(url):
            return jina_fulltext(url, area, outbase)
        return None, {'url': url, 'error': repr(e)}
    ext = infer_ext(url, r, r.content)
    if ext:
        return try_download_document(url, outbase, area, direct_official=is_official(r.url or url))
    html = r.text
    for score, att, txt in attachment_candidates(html, r.url, area):
        p, meta = try_download_document(att, outbase, area, direct_official=is_official(att))
        notes.append({'attachment': att, 'score': score, 'meta': meta})
        if p:
            meta['method'] = 'page_attachment'
            meta['page_url'] = url
            return p, meta
    text, title = page_text(html)
    if allow_full and is_official(r.url or url) and looks_like_full_page(text, title, area):
        pdf = outbase.with_suffix('.pdf')
        ok, detail = print_page_pdf(r.url, pdf)
        if ok:
            valid, meta = validate_document(pdf, area, direct_official=False)
            if valid:
                meta.update({'url': url, 'final_url': r.url, 'method': 'official_html_to_pdf'})
                return pdf, meta
            try:
                pdf.unlink()
            except Exception:
                pass
        p, meta = jina_fulltext(r.url, area, outbase)
        if p:
            return p, meta
    if allow_full and is_official(r.url or url):
        p, meta = jina_fulltext(r.url, area, outbase)
        if p:
            return p, meta
    return None, {'url': url, 'final_url': r.url, 'title': title, 'text_chars': len(clean(text)), 'attempts': notes}


def resolve_so_link(url):
    try:
        r = request(url, timeout=30)
    except Exception:
        return ''
    if 'so.com' not in urllib.parse.urlparse(r.url).netloc:
        return r.url
    html = r.text
    pats = [
        r'URL=[\'\"]?([^\'\"<> ]+)',
        r'location(?:\.href|\.replace)?\s*\(?[\'\"]([^\'\"]+)',
        r'<a[^>]+href=[\'\"]([^\'\"]+)[\'\"]',
    ]
    for pat in pats:
        m = re.search(pat, html, re.I)
        if m:
            u = urllib.parse.unquote(m.group(1).replace('&amp;', '&'))
            if u.startswith('http') and 'so.com' not in urllib.parse.urlparse(u).netloc:
                return u
    return ''


def search_360(area):
    queries = [
        f'"{area}国民经济和社会发展第十五个五年规划纲要"',
        f'"{area}人民政府关于印发{area}国民经济和社会发展第十五个五年规划纲要的通知"',
        f'{area} 第十五个五年规划纲要 pdf',
    ]
    found = []
    seen = set()
    for q in queries:
        try:
            r = request('https://www.so.com/s?q=' + urllib.parse.quote(q), timeout=45)
        except Exception:
            continue
        soup = BeautifulSoup(r.text, 'html.parser')
        for a in soup.select('h3 a, h2 a')[:15]:
            title = a.get_text(' ', strip=True)
            href = a.get('href') or ''
            if not href or not all(k in clean(title) for k in (clean(area), '十五', '规划')):
                continue
            u = resolve_so_link(href) if 'so.com/link' in href else href
            if not u or u in seen:
                continue
            seen.add(u)
            found.append({'query': q, 'title': title, 'url': u})
    return found


def process_target(t, index):
    province, area = t['province'], t['area']
    folder = OUT / safe(province)
    folder.mkdir(parents=True, exist_ok=True)
    filename = safe(province if area == province else province + area)
    outbase = folder / filename
    rec = {'index': index, 'province': province, 'area': area, 'prior_status': t.get('prior_status'), 'attempts': []}
    log(f'[{index:02d}/{len(TARGETS)}] {province} | {area}')

    direct = t.get('direct_url') or ''
    if direct:
        p, meta = try_download_document(direct, outbase, area, direct_official=True)
        rec['attempts'].append({'kind': 'direct', 'meta': meta})
        if p:
            rec.update({'status': 'recovered', 'file': str(p.relative_to(OUT)), 'source_url': direct, 'method': meta.get('method', 'direct_document'), 'validation': meta})
            return rec

    # Process known official full page/seed first.
    for kind, url in [('full_page_seed', t.get('full_page_seed') or ''), ('seed', t.get('seed_url') or '')]:
        if not url:
            continue
        p, meta = process_page(url, outbase, area, allow_full=(kind == 'full_page_seed' or t.get('prior_status','').startswith(('A1','A2'))))
        rec['attempts'].append({'kind': kind, 'meta': meta})
        if p:
            rec.update({'status': 'recovered', 'file': str(p.relative_to(OUT)), 'source_url': meta.get('page_url') or meta.get('url') or url, 'method': meta.get('method', kind), 'validation': meta})
            return rec

    results = search_360(area)
    rec['search_results'] = results
    for item in results:
        u = item['url']
        if not is_official(u):
            continue
        p, meta = process_page(u, outbase, area, allow_full=True)
        rec['attempts'].append({'kind': 'search_result', 'title': item['title'], 'url': u, 'meta': meta})
        if p:
            rec.update({'status': 'recovered', 'file': str(p.relative_to(OUT)), 'source_url': meta.get('page_url') or meta.get('url') or u, 'method': meta.get('method', 'search_result'), 'validation': meta})
            return rec

    rec['status'] = 'not_recovered'
    rec['source_url'] = t.get('seed_url') or ''
    return rec


for i, target in enumerate(TARGETS, 1):
    try:
        rec = process_target(target, i)
    except Exception as e:
        rec = {'index': i, 'province': target['province'], 'area': target['area'], 'status': 'error', 'error': repr(e)}
    REPORT.append(rec)
    log('  => ' + rec.get('status', 'unknown') + (' | ' + rec.get('file','') if rec.get('file') else ''))

(OUT / 'recovery_report.json').write_text(json.dumps(REPORT, ensure_ascii=False, indent=2), encoding='utf-8')
summary = {
    'targets': len(REPORT),
    'recovered': sum(r.get('status') == 'recovered' for r in REPORT),
    'not_recovered': sum(r.get('status') == 'not_recovered' for r in REPORT),
    'errors': sum(r.get('status') == 'error' for r in REPORT),
}
(OUT / 'recovery_summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
log(json.dumps(summary, ensure_ascii=False))
